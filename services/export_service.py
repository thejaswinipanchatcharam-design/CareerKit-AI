"""
CareerKit AI - Document Export Service
Generates professionally formatted PDF and DOCX files for improved resumes
and cover letters using ReportLab and python-docx.
"""

import os
import re
from typing import Dict, Any, Optional
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, ListFlowable, ListItem
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


# Brand Colors Matching Web Interface
PRIMARY_COLOR = colors.HexColor("#6366f1")      # Purple / Indigo Accent
TEXT_DARK = colors.HexColor("#1e293b")          # Slate 800
TEXT_MUTED = colors.HexColor("#64748b")         # Slate 500
BORDER_LIGHT = colors.HexColor("#e2e8f0")       # Slate 200

DOCX_PRIMARY = RGBColor(99, 102, 241)           # #6366f1
DOCX_DARK = RGBColor(30, 41, 59)                # #1e293b
DOCX_MUTED = RGBColor(100, 116, 139)           # #64748b


# ==============================================================================
# REPORTLAB PDF EXPORTS
# ==============================================================================

def export_resume_to_pdf(candidate: Dict[str, Any], resume_text: str, output_path: str) -> str:
    """
    Exports clean, ATS-readable, styled resume PDF using ReportLab.
    """
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()

    # Custom styles
    name_style = ParagraphStyle(
        "CandidateName",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=PRIMARY_COLOR,
        alignment=0,
        spaceAfter=4
    )

    contact_style = ParagraphStyle(
        "ContactInfo",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=TEXT_MUTED,
        spaceAfter=10
    )

    section_heading = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=PRIMARY_COLOR,
        spaceBefore=10,
        spaceAfter=4,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        "ResumeBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13.5,
        textColor=TEXT_DARK,
        spaceAfter=4
    )

    bullet_style = ParagraphStyle(
        "ResumeBullet",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13.5,
        textColor=TEXT_DARK,
        leftIndent=14,
        firstLineIndent=-10,
        spaceAfter=3
    )

    story = []

    # 1. Header Section
    c_name = candidate.get("name", "Candidate Name")
    contact_bits = []
    if candidate.get("email"):
        contact_bits.append(candidate["email"])
    if candidate.get("phone"):
        contact_bits.append(candidate["phone"])
    if candidate.get("location"):
        contact_bits.append(candidate["location"])
    contact_str = " &bull; ".join(contact_bits) if contact_bits else "CareerKit AI Verified Candidate"

    story.append(Paragraph(c_name, name_style))
    story.append(Paragraph(contact_str, contact_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=PRIMARY_COLOR, spaceBefore=0, spaceAfter=8))

    # 2. Parse and flow Markdown/Text Sections
    lines = resume_text.splitlines()
    for line in lines:
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 4))
            continue

        # Format inline bold/italic for ReportLab Paragraph
        formatted = _format_markdown_for_reportlab(stripped)

        # Detect Headings
        if stripped.startswith("#"):
            heading_text = stripped.lstrip("#").strip()
            # If the first heading is the candidate's name, skip it to avoid duplication
            if heading_text.lower() == c_name.lower():
                continue
            story.append(Paragraph(heading_text.upper(), section_heading))
            story.append(HRFlowable(width="100%", thickness=0.75, color=BORDER_LIGHT, spaceBefore=0, spaceAfter=4))
        elif stripped.startswith(("-", "*", "•")):
            bullet_text = formatted.lstrip("-*•").strip()
            story.append(Paragraph(f"&bull; {bullet_text}", bullet_style))
        else:
            story.append(Paragraph(formatted, body_style))

    doc.build(story)
    return output_path


def export_cover_letter_to_pdf(
    candidate: Dict[str, Any],
    letter_text: str,
    output_path: str,
    company: str = "",
    target_role: str = ""
) -> str:
    """
    Exports styled cover letter PDF with proper corporate formatting.
    """
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        rightMargin=48,
        leftMargin=48,
        topMargin=48,
        bottomMargin=48
    )

    styles = getSampleStyleSheet()

    header_name = ParagraphStyle(
        "CLName",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=PRIMARY_COLOR,
        spaceAfter=2
    )

    header_sub = ParagraphStyle(
        "CLSub",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13,
        textColor=TEXT_MUTED,
        spaceAfter=12
    )

    body_style = ParagraphStyle(
        "CLBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=15,
        textColor=TEXT_DARK,
        spaceAfter=10
    )

    story = []

    # Candidate Header
    c_name = candidate.get("name", "Candidate")
    contact_parts = []
    if candidate.get("email"):
        contact_parts.append(candidate["email"])
    if candidate.get("phone"):
        contact_parts.append(candidate["phone"])
    if candidate.get("location"):
        contact_parts.append(candidate["location"])
    contact_str = " | ".join(contact_parts)

    story.append(Paragraph(c_name, header_name))
    if contact_str:
        story.append(Paragraph(contact_str, header_sub))
    story.append(HRFlowable(width="100%", thickness=1, color=PRIMARY_COLOR, spaceBefore=0, spaceAfter=14))

    # Letter body paragraphs
    # Filter out markdown section tags like [COVER LETTER] or [KEY ALIGNMENTS]
    clean_letter = re.sub(r"\[COVER LETTER\]", "", letter_text, flags=re.IGNORECASE)
    parts = re.split(r"\[KEY ALIGNMENTS\]|### Key Alignments", clean_letter, flags=re.IGNORECASE)
    main_letter = parts[0].strip()

    paragraphs = [p.strip() for p in main_letter.split("\n\n") if p.strip()]
    for p in paragraphs:
        formatted = _format_markdown_for_reportlab(p).replace("\n", "<br/>")
        story.append(Paragraph(formatted, body_style))

    doc.build(story)
    return output_path


def _format_markdown_for_reportlab(text: str) -> str:
    """Converts standard markdown bold and italics into ReportLab XML tags."""
    # Escape XML entities
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    # Bold **text** -> <b>text</b>
    text = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", text)
    # Italic *text* -> <i>text</i>
    text = re.sub(r"(?<!\*)\*(?!\*)(.*?)(?<!\*)\*(?!\*)", r"<i>\1</i>", text)
    # Code `code` -> <font face="Courier">\1</font>
    text = re.sub(r"`(.*?)`", r'<font face="Courier">\1</font>', text)
    return text


# ==============================================================================
# PYTHON-DOCX EXPORTS
# ==============================================================================

def export_resume_to_docx(candidate: Dict[str, Any], resume_text: str, output_path: str) -> str:
    """
    Exports clean, structured resume Word document (.docx).
    """
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc = Document()

    # Set 0.75-inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 1. Candidate Name Header
    c_name = candidate.get("name", "Candidate Name")
    name_para = doc.add_paragraph()
    name_run = name_para.add_run(c_name)
    name_run.font.name = "Calibri"
    name_run.font.size = Pt(20)
    name_run.bold = True
    name_run.font.color.rgb = DOCX_PRIMARY
    name_para.paragraph_format.space_after = Pt(2)

    # Contact Info
    contact_parts = []
    if candidate.get("email"):
        contact_parts.append(candidate["email"])
    if candidate.get("phone"):
        contact_parts.append(candidate["phone"])
    if candidate.get("location"):
        contact_parts.append(candidate["location"])
    contact_str = " • ".join(contact_parts) if contact_parts else "CareerKit AI Verified Candidate"

    contact_para = doc.add_paragraph()
    contact_run = contact_para.add_run(contact_str)
    contact_run.font.name = "Calibri"
    contact_run.font.size = Pt(9.5)
    contact_run.font.color.rgb = DOCX_MUTED
    contact_para.paragraph_format.space_after = Pt(12)

    # 2. Body & Sections
    lines = resume_text.splitlines()
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        if stripped.startswith("#"):
            heading_text = stripped.lstrip("#").strip()
            if heading_text.lower() == c_name.lower():
                continue
            h_para = doc.add_paragraph()
            h_run = h_para.add_run(heading_text.upper())
            h_run.font.name = "Calibri"
            h_run.font.size = Pt(12)
            h_run.bold = True
            h_run.font.color.rgb = DOCX_PRIMARY
            h_para.paragraph_format.space_before = Pt(10)
            h_para.paragraph_format.space_after = Pt(3)
        elif stripped.startswith(("-", "*", "•")):
            bullet_text = stripped.lstrip("-*•").strip()
            b_para = doc.add_paragraph(style="List Bullet")
            _add_formatted_runs_docx(b_para, bullet_text)
            b_para.paragraph_format.space_after = Pt(2)
        else:
            p = doc.add_paragraph()
            _add_formatted_runs_docx(p, stripped)
            p.paragraph_format.space_after = Pt(4)

    doc.save(output_path)
    return output_path


def export_cover_letter_to_docx(
    candidate: Dict[str, Any],
    letter_text: str,
    output_path: str,
    company: str = "",
    target_role: str = ""
) -> str:
    """
    Exports professional corporate cover letter Word document (.docx).
    """
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Candidate Name
    c_name = candidate.get("name", "Candidate")
    name_para = doc.add_paragraph()
    name_run = name_para.add_run(c_name)
    name_run.font.name = "Calibri"
    name_run.font.size = Pt(18)
    name_run.bold = True
    name_run.font.color.rgb = DOCX_PRIMARY
    name_para.paragraph_format.space_after = Pt(2)

    # Contact line
    contact_parts = []
    if candidate.get("email"):
        contact_parts.append(candidate["email"])
    if candidate.get("phone"):
        contact_parts.append(candidate["phone"])
    if candidate.get("location"):
        contact_parts.append(candidate["location"])
    contact_str = " | ".join(contact_parts)

    if contact_str:
        sub_para = doc.add_paragraph()
        sub_run = sub_para.add_run(contact_str)
        sub_run.font.name = "Calibri"
        sub_run.font.size = Pt(9.5)
        sub_run.font.color.rgb = DOCX_MUTED
        sub_para.paragraph_format.space_after = Pt(16)

    # Clean letter body
    clean_letter = re.sub(r"\[COVER LETTER\]", "", letter_text, flags=re.IGNORECASE)
    parts = re.split(r"\[KEY ALIGNMENTS\]|### Key Alignments", clean_letter, flags=re.IGNORECASE)
    main_letter = parts[0].strip()

    paragraphs = [p.strip() for p in main_letter.split("\n\n") if p.strip()]
    for p_text in paragraphs:
        p = doc.add_paragraph()
        _add_formatted_runs_docx(p, p_text)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15

    doc.save(output_path)
    return output_path


def _add_formatted_runs_docx(paragraph, text: str) -> None:
    """Parses **bold** markdown segments into python-docx formatted runs."""
    # Split by bold markdown
    tokens = re.split(r"(\*\*.*?\*\*)", text)
    for token in tokens:
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.font.color.rgb = DOCX_DARK
        else:
            if token:
                run = paragraph.add_run(token)
                run.font.name = "Calibri"
                run.font.size = Pt(10)
                run.font.color.rgb = DOCX_DARK
