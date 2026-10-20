"""
Tests for Document Parsing (PDF, DOCX, TXT), Ingestion Routes, and AI Service
"""

import io
from werkzeug.datastructures import FileStorage
from services.resume_parser import extract_text_from_file, is_allowed_file
from services.ai_service import generate_ai_response
from config import Config


def test_is_allowed_file():
    """Validates allowed and disallowed file extensions."""
    assert is_allowed_file("resume.pdf") is True
    assert is_allowed_file("cv.docx") is True
    assert is_allowed_file("resume.txt") is True
    assert is_allowed_file("resume.PDF") is True  # Case insensitive
    assert is_allowed_file("script.py") is False
    assert is_allowed_file("image.png") is False
    assert is_allowed_file("virus.exe") is False
    assert is_allowed_file("") is False
    assert is_allowed_file(None) is False


def test_txt_parsing():
    """Tests text extraction from plain text files."""
    sample_text = "John Doe\nEducation: B.Tech Computer Science\nSkills: Python, Flask"
    file_bytes = io.BytesIO(sample_text.encode("utf-8"))
    storage = FileStorage(stream=file_bytes, filename="my_resume.txt", content_type="text/plain")

    text, err, f_type = extract_text_from_file(storage)
    assert err is None
    assert f_type == "txt"
    assert "John Doe" in text
    assert "Python, Flask" in text


def test_docx_parsing():
    """Generates an in-memory DOCX and verifies paragraph and table text extraction."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Alex Rivera - Senior Full Stack Developer")
    doc.add_paragraph("Education: Master of Science in Software Engineering")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skill Category"
    table.rows[0].cells[1].text = "Python, SQLite, Docker"

    doc_stream = io.BytesIO()
    doc.save(doc_stream)
    doc_stream.seek(0)

    storage = FileStorage(
        stream=doc_stream,
        filename="rivera_resume.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    text, err, f_type = extract_text_from_file(storage)
    assert err is None
    assert f_type == "docx"
    assert "Alex Rivera" in text
    assert "Software Engineering" in text
    assert "Python, SQLite, Docker" in text


def test_pdf_parsing():
    """Generates an in-memory text-based PDF using ReportLab and tests extraction with pypdf."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    pdf_stream = io.BytesIO()
    c = canvas.Canvas(pdf_stream, pagesize=letter)
    c.drawString(100, 750, "Priya Patel - Software Engineer")
    c.drawString(100, 720, "Education: B.E. Computer Science 2026")
    c.drawString(100, 690, "Technical Skills: Flask, PostgreSQL, React, AWS")
    c.showPage()
    c.save()
    pdf_stream.seek(0)

    storage = FileStorage(
        stream=pdf_stream,
        filename="patel_resume.pdf",
        content_type="application/pdf"
    )

    text, err, f_type = extract_text_from_file(storage)
    assert err is None
    assert f_type == "pdf"
    assert "Priya Patel" in text
    assert "Software Engineer" in text
    assert "Flask, PostgreSQL" in text


def test_empty_file_rejected():
    """Verifies that 0-byte empty files are rejected."""
    empty_stream = io.BytesIO(b"")
    storage = FileStorage(stream=empty_stream, filename="empty_file.txt", content_type="text/plain")

    text, err, f_type = extract_text_from_file(storage)
    assert text is None
    assert "empty" in err.lower()


def test_unsupported_file_rejected():
    """Verifies that unapproved extensions are rejected."""
    dummy_stream = io.BytesIO(b"malicious executable or binary")
    storage = FileStorage(stream=dummy_stream, filename="payload.exe", content_type="application/x-msdownload")

    text, err, f_type = extract_text_from_file(storage)
    assert text is None
    assert "unsupported" in err.lower()


def test_manual_resume_ingest_route(client):
    """Tests ingesting resume content via manual text submission."""
    # First create a candidate
    client.post("/candidates/new", data={
        "name": "Devin Miller",
        "education": "B.Tech IT",
        "skills": "Python, Django",
        "target_role": "Backend Engineer"
    }, follow_redirects=True)

    # Ingest manual resume
    response = client.post("/resume/ingest", data={
        "candidate_id": 1,
        "input_method": "manual",
        "manual_resume_text": "Devin Miller\nExperience: 1 year backend developer\nProjects: REST API Gateway",
        "target_job_title": "Backend Engineer",
        "target_company": "CloudTech",
        "job_description": "Seeking Python and Django engineer."
    }, follow_redirects=True)

    assert response.status_code == 200
    assert b"successfully ingested" in response.data
    assert b"Devin Miller" in response.data
    assert b"REST API Gateway" in response.data


def test_file_upload_resume_ingest_route(client):
    """Tests uploading a text file resume via the web route."""
    # Create candidate
    client.post("/candidates/new", data={
        "name": "Sara Lin",
        "education": "B.S. Software Eng",
        "skills": "Go, Python",
        "target_role": "Cloud Developer"
    }, follow_redirects=True)

    file_bytes = io.BytesIO(b"Sara Lin\nCloud Developer\nEducation: B.S. Software Eng\nSkills: Go, Docker, Kubernetes")

    response = client.post("/resume/ingest", data={
        "candidate_id": 1,
        "input_method": "upload",
        "resume_file": (file_bytes, "sara_lin_resume.txt"),
        "target_job_title": "Cloud Developer",
        "target_company": "FinTech Corp",
        "job_description": "Experience with Go and Docker required."
    }, content_type="multipart/form-data", follow_redirects=True)

    assert response.status_code == 200
    assert b"successfully ingested" in response.data
    assert b"Sara Lin" in response.data
    assert b"Kubernetes" in response.data


def test_ai_service_unified_interface():
    """Tests the unified AI service interface and offline mock fallback."""
    # Empty prompt check
    res_empty = generate_ai_response("")
    assert res_empty["success"] is False
    assert res_empty["error_type"] == "EMPTY_PROMPT"

    # Non-empty prompt in test environment (defaults to structured mock response if no key)
    res_valid = generate_ai_response("Please improve this candidate profile for software engineering.")
    assert "success" in res_valid
    assert "content" in res_valid
    assert "provider" in res_valid
    assert "model" in res_valid
    assert res_valid["success"] is True
    assert "Professional Summary" in res_valid["content"]
    assert "Key Improvements" in res_valid["content"]
